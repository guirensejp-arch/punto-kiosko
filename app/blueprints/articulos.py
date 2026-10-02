from flask import Blueprint, flash, redirect, render_template, request, url_for
from flask_login import login_required
from openpyxl import Workbook

from app.decorators import admin_required
from app.extensions import db
from app.forms import ProductoForm, ProductoInsumoForm
from app.models.proveedor import Insumo
from app.models.articulo import Articulo, ArticuloInsumo, CodigoBarras
from app.services.food_cost import costo_producto, margen_producto
from app.utils.auditoria import registrar
from app.utils.excel import (
    FORMATO_MONEDA,
    ajustar_anchos,
    formatear_columna,
    marcar_encabezado,
    pesos,
    respuesta_xlsx,
)
from app.utils.moneda import centavos_a_editable, parsear_centavos
from app.utils.numeros import parsear_decimal
from app.utils.unidades import convertir

articulos_bp = Blueprint('articulos', __name__, url_prefix='/articulos')


def _insumos_activos():
    return (
        Insumo.query.filter_by(activo=True).order_by(Insumo.nombre).all()
    )


@articulos_bp.route('/')
@login_required
def lista():
    q = (request.args.get('q') or '').strip()
    categoria = (request.args.get('categoria') or '').strip()

    consulta = Articulo.query
    if q:
        consulta = consulta.filter(Articulo.nombre.ilike(f'%{q}%'))
    if categoria:
        consulta = consulta.filter(Articulo.categoria == categoria)

    articulos = consulta.order_by(
        Articulo.activo.desc(), Articulo.categoria, Articulo.nombre
    ).all()
    categorias = [
        c[0]
        for c in db.session.query(Articulo.categoria)
        .filter(Articulo.categoria.isnot(None))
        .distinct()
        .order_by(Articulo.categoria)
        .all()
    ]

    filas = [
        {
            'articulo': p,
            'costo': costo_producto(p),
            'margen': margen_producto(p),
        }
        for p in articulos
    ]
    return render_template(
        'articulos/lista.html',
        filas=filas,
        categorias=categorias,
        q=q,
        categoria=categoria,
    )


@articulos_bp.route('/nuevo', methods=['GET', 'POST'])
@login_required
@admin_required
def nuevo():
    form = ProductoForm()
    if form.validate_on_submit():
        try:
            precio = parsear_centavos(form.precio_venta.data)
        except ValueError:
            flash('El precio no es un monto válido.', 'danger')
            return render_template('articulos/nuevo.html', form=form)

        articulo = Articulo(
            nombre=form.nombre.data.strip(),
            descripcion=(form.descripcion.data or '').strip() or None,
            categoria=form.categoria.data or None,
            precio_venta=precio,
            margen_objetivo=form.margen_objetivo.data,
            activo=form.activo.data,
        )
        db.session.add(articulo)
        db.session.commit()
        registrar('CREAR_PRODUCTO', 'articulo', articulo.id)
        flash('Articulo creado. Ahora cargá los insumos de la receta.', 'success')
        return redirect(url_for('articulos.detalle', articulo_id=articulo.id))

    return render_template('articulos/nuevo.html', form=form)


@articulos_bp.route('/<int:articulo_id>')
@login_required
def detalle(articulo_id):
    articulo = db.get_or_404(Articulo, articulo_id)
    form = ProductoForm(obj=articulo)
    form.precio_venta.data = centavos_a_editable(articulo.precio_venta)

    linea_form = ProductoInsumoForm()
    linea_form.insumo_id.choices = [(i.id, f'{i.nombre} ({i.unidad})') for i in _insumos_activos()]

    return render_template(
        'articulos/detalle.html',
        articulo=articulo,
        form=form,
        linea_form=linea_form,
        costo=costo_producto(articulo),
        margen=margen_producto(articulo),
    )


@articulos_bp.route('/<int:articulo_id>/guardar', methods=['POST'])
@login_required
@admin_required
def guardar(articulo_id):
    articulo = db.get_or_404(Articulo, articulo_id)
    form = ProductoForm()

    if form.validate_on_submit():
        try:
            articulo.precio_venta = parsear_centavos(form.precio_venta.data)
        except ValueError:
            flash('El precio no es un monto válido.', 'danger')
            return redirect(url_for('articulos.detalle', articulo_id=articulo.id))

        articulo.nombre = form.nombre.data.strip()
        articulo.descripcion = (form.descripcion.data or '').strip() or None
        articulo.categoria = form.categoria.data or None
        articulo.margen_objetivo = form.margen_objetivo.data
        articulo.activo = form.activo.data
        db.session.commit()
        registrar('EDITAR_PRODUCTO', 'articulo', articulo.id)
        flash('Articulo actualizado.', 'success')
    else:
        flash('Revisá los datos del articulo.', 'danger')

    return redirect(url_for('articulos.detalle', articulo_id=articulo.id))


@articulos_bp.route('/<int:articulo_id>/insumos', methods=['POST'])
@login_required
@admin_required
def agregar_insumo(articulo_id):
    articulo = db.get_or_404(Articulo, articulo_id)
    form = ProductoInsumoForm()
    form.insumo_id.choices = [(i.id, i.nombre) for i in _insumos_activos()]

    if not form.validate_on_submit():
        flash('Revisá los datos del insumo.', 'danger')
        return redirect(url_for('articulos.detalle', articulo_id=articulo.id))

    insumo = db.session.get(Insumo, form.insumo_id.data)
    if insumo is None or not insumo.activo:
        flash('El insumo no existe o está inactivo.', 'danger')
        return redirect(url_for('articulos.detalle', articulo_id=articulo.id))

    try:
        cantidad = parsear_decimal(form.cantidad.data)
    except ValueError:
        cantidad = None
    if cantidad is None or cantidad <= 0:
        flash('La cantidad debe ser un número mayor a 0.', 'danger')
        return redirect(url_for('articulos.detalle', articulo_id=articulo.id))

    if convertir(cantidad, form.unidad.data, insumo.unidad) is None:
        flash(
            f'La unidad {form.unidad.data} no es compatible con la unidad del '
            f'insumo ({insumo.unidad}).',
            'danger',
        )
        return redirect(url_for('articulos.detalle', articulo_id=articulo.id))

    existente = ArticuloInsumo.query.filter_by(
        articulo_id=articulo.id, insumo_id=insumo.id
    ).first()
    if existente:
        existente.cantidad = cantidad
        existente.unidad = form.unidad.data
        flash('El insumo ya estaba en la receta: se actualizó la cantidad.', 'info')
    else:
        db.session.add(
            ArticuloInsumo(
                articulo_id=articulo.id,
                insumo_id=insumo.id,
                cantidad=cantidad,
                unidad=form.unidad.data,
            )
        )
        flash('Insumo agregado a la receta.', 'success')

    db.session.commit()
    registrar('EDITAR_RECETA', 'articulo', articulo.id, {'insumo_id': insumo.id})
    return redirect(url_for('articulos.detalle', articulo_id=articulo.id))


@articulos_bp.route('/<int:articulo_id>/insumos/<int:linea_id>/quitar', methods=['POST'])
@login_required
@admin_required
def quitar_insumo(articulo_id, linea_id):
    linea = db.get_or_404(ArticuloInsumo, linea_id)
    if linea.articulo_id != articulo_id:
        flash('La línea no pertenece a ese articulo.', 'danger')
        return redirect(url_for('articulos.detalle', articulo_id=articulo_id))

    db.session.delete(linea)
    db.session.commit()
    registrar('EDITAR_RECETA', 'articulo', articulo_id, {'quitar_insumo': linea.insumo_id})
    flash('Insumo quitado de la receta.', 'info')
    return redirect(url_for('articulos.detalle', articulo_id=articulo_id))


@articulos_bp.route('/<int:articulo_id>/desactivar', methods=['POST'])
@login_required
@admin_required
def desactivar(articulo_id):
    articulo = db.get_or_404(Articulo, articulo_id)
    articulo.activo = False
    db.session.commit()
    registrar('DESACTIVAR_PRODUCTO', 'articulo', articulo.id)
    flash(f'Articulo {articulo.nombre} desactivado.', 'info')
    return redirect(url_for('articulos.lista'))


@articulos_bp.route('/<int:articulo_id>/activar', methods=['POST'])
@login_required
@admin_required
def activar(articulo_id):
    articulo = db.get_or_404(Articulo, articulo_id)
    articulo.activo = True
    db.session.commit()
    registrar('ACTIVAR_PRODUCTO', 'articulo', articulo.id)
    flash(f'Articulo {articulo.nombre} activado.', 'success')
    return redirect(url_for('articulos.lista'))


@articulos_bp.route('/exportar.xlsx')
@login_required
def exportar_xlsx():
    articulos = Articulo.query.order_by(
        Articulo.categoria, Articulo.nombre
    ).all()

    libro = Workbook()
    hoja = libro.active
    hoja.title = 'Recetas'
    hoja.append(
        ['Articulo', 'Categoría', 'Precio venta', 'Costo insumos', 'Margen %', 'Estado']
    )
    for articulo in articulos:
        margen = margen_producto(articulo)
        hoja.append([
            articulo.nombre,
            articulo.categoria or '',
            pesos(articulo.precio_venta),
            pesos(costo_producto(articulo)),
            margen if margen is not None else '',
            'Activo' if articulo.activo else 'Inactivo',
        ])

    marcar_encabezado(hoja)
    formatear_columna(hoja, 3, FORMATO_MONEDA)
    formatear_columna(hoja, 4, FORMATO_MONEDA)
    ajustar_anchos(hoja)
    hoja.freeze_panes = 'A2'

    return respuesta_xlsx(libro, 'articulos.xlsx')


def _preview_remarcar(articulos, modo, valor):
    """Calcula precio nuevo por artículo sin guardar. Devuelve (filas, con_cambio)."""
    filas = []
    con_cambio = 0
    for art in articulos:
        if modo == 'PORCENTAJE':
            nuevo = int(round(art.precio_venta * (1 + valor / 100)))
        else:  # MONTO: suma fija de valor (puede ser negativa)
            nuevo = art.precio_venta + valor
        if nuevo < 0:
            nuevo = 0
        if nuevo != art.precio_venta:
            con_cambio += 1
        filas.append({'articulo': art, 'nuevo': nuevo})
    return filas, con_cambio


@articulos_bp.route('/remarcar', methods=['GET', 'POST'])
@login_required
@admin_required
def remarcar():
    """Remarcado masivo de precios por categoría o para todos.

    Modos: porcentaje (+/-) o monto fijo (+/-) en pesos. Muestra una previsualización
    y aplica recién al confirmar.
    """
    categoria = (request.args.get('categoria') or request.form.get('categoria') or '').strip()
    modo = (request.form.get('modo') or request.args.get('modo') or 'PORCENTAJE').strip()
    valor_txt = (request.form.get('valor') or request.args.get('valor') or '').strip()
    aplicar = request.form.get('aplicar') == '1'

    categorias = [
        c[0]
        for c in db.session.query(Articulo.categoria)
        .filter(Articulo.activo.is_(True), Articulo.categoria.isnot(None))
        .distinct()
        .order_by(Articulo.categoria)
        .all()
    ]

    articulos = None
    filas = []
    con_cambio = 0
    valor = None

    if valor_txt:
        if modo == 'PORCENTAJE':
            valor = parsear_decimal(valor_txt)
        else:
            valor = parsear_centavos(valor_txt)
        if valor is None:
            flash('Valor inválido.', 'danger')
        else:
            consulta = Articulo.query.filter_by(activo=True)
            if categoria:
                consulta = consulta.filter(Articulo.categoria == categoria)
            articulos = consulta.order_by(Articulo.categoria, Articulo.nombre).all()
            filas, con_cambio = _preview_remarcar(articulos, modo, valor)

            if aplicar:
                for fila in filas:
                    fila['articulo'].precio_venta = fila['nuevo']
                db.session.commit()
                registrar(
                    'REMARCAR_PRECIOS', 'articulo', None,
                    {'modo': modo, 'valor': float(valor), 'categoria': categoria or 'TODAS',
                     'artefactos': len(filas)},
                )
                flash(f'Precios actualizados: {len(articulos)} artículos.', 'success')
                return redirect(url_for('articulos.lista'))

    return render_template(
        'articulos/remarcar.html',
        categorias=categorias,
        categoria=categoria,
        modo=modo,
        valor_txt=valor_txt,
        filas=filas,
        con_cambio=con_cambio,
        aplicado=aplicar and articulos is not None,
    )


COLUMNAS_IMPORT = [
    'nombre', 'categoria', 'sku', 'codigo_barras',
    'precio_venta', 'precio_costo', 'stock', 'unidad', 'stock_propio',
]


def _leer_filas_import(archivo):
    """Lee CSV o XLSX y devuelve lista de dicts con las columnas esperadas."""
    nombre = (archivo.filename or '').lower()
    filas = []
    if nombre.endswith('.xlsx'):
        from openpyxl import load_workbook

        libro = load_workbook(archivo, read_only=True, data_only=True)
        hoja = libro.active
        iterador = hoja.iter_rows(values_only=True)
        encabezados = None
        for crudo in iterador:
            if encabezados is None:
                encabezados = [str(c).strip().lower() if c is not None else '' for c in crudo]
                continue
            filas.append(dict(zip(encabezados, crudo)))
        libro.close()
    else:
        import csv
        import io

        texto = archivo.read().decode('utf-8-sig')
        lector = csv.DictReader(io.StringIO(texto))
        for crudo in lector:
            filas.append({(k or '').strip().lower(): v for k, v in crudo.items()})
    return filas


@articulos_bp.route('/importar', methods=['GET', 'POST'])
@login_required
@admin_required
def importar():
    """Importa/actualiza artículos desde CSV o XLSX.

    Columnas: nombre, categoria, sku, codigo_barras, precio_venta, precio_costo,
    stock, unidad, stock_propio. Se actualiza por SKU o código de barras si existe.
    """
    resultado = None
    if request.method == 'POST':
        archivo = request.files.get('archivo')
        if archivo is None or not archivo.filename:
            flash('Elegí un archivo CSV o XLSX.', 'danger')
            return redirect(url_for('articulos.importar'))
        try:
            crudas = _leer_filas_import(archivo)
        except Exception as error:  # noqa: BLE001 - mostrar el error al usuario
            flash(f'No se pudo leer el archivo: {error}', 'danger')
            return redirect(url_for('articulos.importar'))

        creados = actualizados = errores = 0
        for cruda in crudas:
            nombre = str(cruda.get('nombre') or '').strip()
            if not nombre:
                errores += 1
                continue
            sku = str(cruda.get('sku') or '').strip() or None
            codigo = str(cruda.get('codigo_barras') or '').strip() or None

            articulo = None
            if sku:
                articulo = Articulo.query.filter_by(sku=sku).first()
            if articulo is None and codigo:
                cb = CodigoBarras.query.filter_by(codigo=codigo).first()
                if cb is not None:
                    articulo = cb.articulo

            try:
                precio = parsear_centavos(str(cruda.get('precio_venta') or '0'))
            except ValueError:
                errores += 1
                continue
            try:
                costo = parsear_centavos(str(cruda.get('precio_costo') or '0'))
            except ValueError:
                costo = 0
            stock_txt = str(cruda.get('stock') or '').strip()
            unidad = str(cruda.get('unidad') or 'ud').strip() or 'ud'
            propio = str(cruda.get('stock_propio') or '').strip().lower() in ('1', 'si', 'sí', 'true', 'x')

            if articulo is None:
                articulo = Articulo(nombre=nombre, precio_venta=precio)
                db.session.add(articulo)
                creados += 1
            else:
                actualizados += 1
            articulo.nombre = nombre
            articulo.categoria = str(cruda.get('categoria') or '').strip() or None
            articulo.sku = sku
            articulo.precio_venta = precio
            articulo.precio_costo = costo
            if stock_txt:
                valor_stock = parsear_decimal(stock_txt)
                if valor_stock is not None:
                    articulo.stock = valor_stock
            articulo.unidad = unidad
            if propio:
                articulo.stock_propio = True
            db.session.flush()

            if codigo and CodigoBarras.query.filter_by(codigo=codigo).first() is None:
                db.session.add(CodigoBarras(codigo=codigo, articulo_id=articulo.id))

        db.session.commit()
        registrar('IMPORTAR_ARTICULOS', 'articulo', None,
                  {'creados': creados, 'actualizados': actualizados, 'errores': errores})
        resultado = {'creados': creados, 'actualizados': actualizados, 'errores': errores}
        flash(f'Importación: {creados} creados, {actualizados} actualizados, {errores} con error.', 'success')

    return render_template('articulos/importar.html', columnas=COLUMNAS_IMPORT, resultado=resultado)
